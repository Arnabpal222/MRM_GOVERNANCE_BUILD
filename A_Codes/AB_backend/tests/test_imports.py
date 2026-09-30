"""Import engine tests (BRD §41–§47, IMP-01..09; design TC-IMP-*)."""
import csv
import io

from openpyxl import load_workbook

from tests.conftest import ADMIN, AUDITOR, VALIDATOR, login

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _upload(client, headers, files, **form):
    """files: list of (name, bytes)."""
    multipart = [("files", (name, data, XLSX if name.endswith(".xlsx") else "text/csv")) for name, data in files]
    return client.post("/api/imports/batches", headers=headers, files=multipart, data=form)


def _issues(client, headers, batch_id):
    return client.get(f"/api/imports/batches/{batch_id}/issues", headers=headers).json()


def _csv(rows: list[dict]) -> bytes:
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=list(rows[0]))
    w.writeheader()
    w.writerows(rows)
    return buf.getvalue().encode()


def test_template_registry_listing(client):
    t = {x["template_id"]: x for x in client.get("/api/imports/templates", headers=login(client, *ADMIN)).json()}
    assert all(t[k]["available"] for k in ("T12", "T01", "T02", "T03", "T04", "T05"))
    assert not t["T06"]["available"] and t["T06"]["phase"] == "Phase 7"
    assert [x for x in sorted(t.values(), key=lambda x: x["order"]) if x["available"]][0]["template_id"] == "T12"
    assert t["T02"]["prerequisites_met"] is True  # bootstrap users exist
    assert t["T03"]["prerequisites_met"] is False and t["T03"]["missing_prerequisites"] == ["T02"]
    model_type = next(c for c in t["T02"]["columns"] if c["name"] == "model_type")
    assert model_type["allowed"][0] == "Credit Risk"  # allowed values come from policy


def test_tc_imp_06_blank_templates_match_registry(client):
    h = login(client, *ADMIN)
    templates = client.get("/api/imports/templates", headers=h).json()
    for t in (x for x in templates if x["available"]):
        resp = client.get(f"/api/imports/templates/{t['template_id']}/blank", headers=h)
        assert resp.status_code == 200
        wb = load_workbook(io.BytesIO(resp.content))
        assert wb.sheetnames == ["data", "instructions"]
        assert [c.value for c in wb["data"][1]] == [c["name"] for c in t["columns"]]
    as_csv = client.get("/api/imports/templates/T04/blank?format=csv", headers=h)
    assert as_csv.text.strip().split(",")[0] == "finding_id"


def test_tc_usr_03_only_importers_can_upload(client, demo_dir):
    data = (demo_dir / "T02_new_models_demo.xlsx").read_bytes()
    for who in (AUDITOR, VALIDATOR):
        assert _upload(client, login(client, *who), [("T02_new_models_demo.xlsx", data)]).status_code == 403


def test_tc_imp_03_1_preview_then_partial_load(client, demo_dir):
    h = login(client, *ADMIN)
    resp = _upload(client, h, [("T02_new_models_demo.xlsx", (demo_dir / "T02_new_models_demo.xlsx").read_bytes())])
    assert resp.status_code == 201, resp.text
    batch = resp.json()
    assert batch["batch_id"] == "B-20260930-001" and batch["status"] == "Ready"
    f = batch["files"][0]
    assert (f["template_id"], f["template_source"]) == ("T02", "filename")
    assert (f["rows_total"], f["rows_valid"], f["rows_invalid"], f["rows_new"]) == (5, 4, 1, 4)
    assert len(f["checksum"]) == 64
    # TC-IMP-01-1: nothing is saved until Load
    assert client.get("/api/models", headers=h).json() == []
    errors = [i for i in _issues(client, h, batch["batch_id"]) if i["severity"] == "error"]
    assert len(errors) == 1
    assert (errors[0]["row_number"], errors[0]["column_name"]) == (5, "validator_id")
    assert "cannot validate model M-0304 because U-007 is also the model owner" in errors[0]["message"]

    loaded = client.post(f"/api/imports/batches/{batch['batch_id']}/load", headers=h).json()
    assert loaded["status"] == "Partially Completed"
    assert (loaded["successful_records"], loaded["failed_records"]) == (4, 1)
    models = {m["model_id"]: m for m in client.get("/api/models", headers=h).json()}
    assert set(models) == {"M-0301", "M-0302", "M-0303", "M-0305"}
    # TC-TIR-01-3: tiers calculated on load
    assert [models[m]["effective_tier"] for m in ("M-0301", "M-0302", "M-0303", "M-0305")] == \
        ["High", "High", "Medium", "Medium"]
    # lineage: every created record's audit event carries the batch id
    events = client.get("/api/audit", headers=h, params={"batch_id": batch["batch_id"], "entity": "model"}).json()
    assert {e["entity_id"] for e in events} == set(models)

    report = client.get(f"/api/imports/batches/{batch['batch_id']}/error-report", headers=h)
    lines = list(csv.reader(io.StringIO(report.content.decode("utf-8-sig"))))
    assert lines[0] == ["file", "sheet", "row", "column", "severity", "error_type", "message", "original_value"]
    assert lines[1][:4] == ["T02_new_models_demo.xlsx", "data", "5", "validator_id"] and lines[1][7] == "U-007"


def test_tc_imp_05_reimport_updates_and_is_idempotent(client, demo_dir):
    h = login(client, *ADMIN)
    data = (demo_dir / "T02_new_models_demo.xlsx").read_bytes()
    b1 = _upload(client, h, [("T02_new_models_demo.xlsx", data)]).json()
    client.post(f"/api/imports/batches/{b1['batch_id']}/load", headers=h)

    again = _upload(client, h, [("T02_new_models_demo.xlsx", data)]).json()
    f = again["files"][0]
    assert (f["rows_new"], f["rows_update"], f["rows_unchanged"]) == (0, 0, 4)
    assert f["duplicate_of_batch"] == b1["batch_id"]  # identical file already loaded

    wb = load_workbook(io.BytesIO(data))
    ws = wb["data"]
    col = [c.value for c in ws[1]].index("version") + 1
    ws.cell(row=2, column=col, value="v1.1")  # M-0301
    buf = io.BytesIO()
    wb.save(buf)
    b3 = _upload(client, h, [("T02_new_models_demo.xlsx", buf.getvalue())]).json()
    assert b3["files"][0]["rows_update"] == 1
    client.post(f"/api/imports/batches/{b3['batch_id']}/load", headers=h)
    events = client.get("/api/audit", headers=h, params={"batch_id": b3["batch_id"], "action": "update"}).json()
    ev = next(e for e in events if e["entity_id"] == "M-0301")
    assert (ev["before_json"]["version"], ev["after_json"]["version"]) == ("v1.0", "v1.1")


def test_tc_imp_01_2_wrong_file_for_template(client):
    h = login(client, *ADMIN)
    rows = [{"finding_id": "F-0001", "model_id": "M-0012", "title": "x"}]
    batch = _upload(client, h, [("T02_models.csv", _csv(rows))]).json()
    f = batch["files"][0]
    assert f["status"] == "Invalid" and "Required column(s) missing" in f["file_error"]
    assert batch["status"] == "Validation Failed"


def test_tc_imp_01_3_non_spreadsheet_rejected(client):
    resp = _upload(client, login(client, *ADMIN), [("report.pdf", b"%PDF-1.7")])
    assert resp.status_code == 422 and "only .xlsx and .csv" in resp.json()["detail"]


def test_path_traversal_rejected(client):
    resp = _upload(client, login(client, *ADMIN), [("x.csv", b"a,b\n1,2\n")], paths=["../../etc/T01.csv"])
    assert resp.status_code == 422


def test_row_level_errors_by_type(loaded_client, demo_dir):
    h = login(loaded_client, *ADMIN)
    name = "T04_findings_demo_errors.csv"
    batch = _upload(loaded_client, h, [(name, (demo_dir / name).read_bytes())]).json()
    f = batch["files"][0]
    assert (f["rows_valid"], f["rows_invalid"]) == (1, 4)
    by_row = {i["row_number"]: i for i in _issues(loaded_client, h, batch["batch_id"]) if i["severity"] == "error"}
    assert by_row[3]["error_type"] == "allowed" and "Allowed: Critical, Medium, Low" in by_row[3]["message"]
    assert by_row[4]["error_type"] == "reference" and by_row[4]["message"] == "Model M-9999 does not exist."
    assert by_row[5]["error_type"] == "rule" and "Closed date is required" in by_row[5]["message"]
    assert by_row[6]["error_type"] == "date" and by_row[6]["original_value"] == "15/09/2026"


def test_tc_imp_02_4_duplicate_keys_in_file(client):
    h = login(client, *ADMIN)
    base = {"user_id": "U-500", "full_name": "A Person", "email": "a@demo-bank.example", "role": "Validator",
            "business_line": "", "active": "Y"}
    rows = [base, {**base, "email": "b@demo-bank.example"}]
    batch = _upload(client, h, [("T01_users.csv", _csv(rows))]).json()
    errors = [i for i in _issues(client, h, batch["batch_id"]) if i["error_type"] == "duplicate_key"]
    assert [e["row_number"] for e in errors] == [2, 3]


def test_cross_file_dependencies_in_one_batch(client, demo_dir):
    """Findings referencing models created by another file in the same batch validate and load (IMP-06)."""
    h = login(client, *ADMIN)
    finding = {"finding_id": "F-0100", "model_id": "M-0301", "validation_id": "", "title": "Sample gap",
               "description": "Sample excludes 2025 vintages.", "severity": "Medium", "category": "Data",
               "status": "Open", "owner_id": "U-003", "raised_date": "2026-09-01", "due_date": "2026-12-31",
               "closed_date": ""}
    files = [("T04_new_findings.csv", _csv([finding])),
             ("T02_new_models_demo.xlsx", (demo_dir / "T02_new_models_demo.xlsx").read_bytes())]
    batch = _upload(client, h, files).json()
    assert {f["template_id"]: f["rows_valid"] for f in batch["files"]} == {"T04": 1, "T02": 4}
    loaded = client.post(f"/api/imports/batches/{batch['batch_id']}/load", headers=h).json()
    assert loaded["successful_records"] == 5
    assert client.get("/api/findings", headers=h, params={"model_id": "M-0301"}).json()[0]["source"] == "Import"


def test_load_requires_validated_batch_and_cancel(client):
    h = login(client, *ADMIN)
    rows = [{"finding_id": "F-0001", "model_id": "M-0012"}]
    batch = _upload(client, h, [("T02_x.csv", _csv(rows))]).json()
    assert batch["status"] == "Validation Failed"
    assert client.post(f"/api/imports/batches/{batch['batch_id']}/load", headers=h).status_code == 409
    cancelled = client.post(f"/api/imports/batches/{batch['batch_id']}/cancel", headers=h).json()
    assert cancelled["status"] == "Cancelled"
    history = client.get("/api/imports/batches", headers=h).json()
    assert history[0]["batch_id"] == batch["batch_id"]


def test_user_can_remap_file_template(client, demo_dir):
    h = login(client, *ADMIN)
    data = (demo_dir / "T02_new_models_demo.xlsx").read_bytes()
    batch = _upload(client, h, [("inventory_update.xlsx", data)]).json()
    f = batch["files"][0]
    assert (f["template_id"], f["template_source"]) == ("T02", "header")  # detected from the columns
    resp = client.patch(f"/api/imports/files/{f['import_file_id']}", headers=h, json={"template_id": "T03"})
    assert resp.json()["template_source"] == "user"
    again = client.post(f"/api/imports/batches/{batch['batch_id']}/validate", headers=h).json()
    assert again["files"][0]["status"] == "Invalid"


def test_on_existing_reject(client, demo_dir):
    h = login(client, *ADMIN)
    data = (demo_dir / "T02_new_models_demo.xlsx").read_bytes()
    b1 = _upload(client, h, [("T02_new_models_demo.xlsx", data)]).json()
    client.post(f"/api/imports/batches/{b1['batch_id']}/load", headers=h)
    b2 = _upload(client, h, [("T02_new_models_demo.xlsx", data)], on_existing="reject").json()
    assert b2["files"][0]["rows_valid"] == 0 and b2["status"] == "Validation Failed"


def test_tc_gen_01_reset_to_seed(client, test_settings):
    h = login(client, *ADMIN)
    assert client.post("/api/admin/reset", headers=login(client, *AUDITOR)).status_code == 403
    resp = client.post("/api/admin/reset", headers=h)
    assert resp.status_code == 202, resp.text
    batch = client.get(f"/api/imports/batches/{resp.json()['batch_id']}", headers=login(client, *ADMIN)).json()
    assert batch["status"] == "Completed", batch
    assert batch["source"] == "reset" and batch["failed_records"] == 0
    models = client.get("/api/models", headers=login(client, *ADMIN)).json()
    assert len(models) == 256
    tiers = {t: sum(1 for m in models if m["effective_tier"] == t) for t in ("High", "Medium", "Low")}
    assert tiers == {"High": 89, "Medium": 106, "Low": 61}
    column = [m for m in models if m["state"]["displayed_column"] == "Revalidation"]
    assert len(column) == 31 and sum(1 for m in column if m["state"]["revalidation_status"] == "Overdue") == 8
    m78 = next(m for m in models if m["model_id"] == "M-0078")
    assert m78["state"]["sod_breach"]  # legacy conflict loaded, flagged
    # the audit trail survives the reset and records it
    resets = client.get("/api/audit", headers=login(client, *ADMIN), params={"action": "reset"}).json()
    assert len(resets) == 1 and resets[0]["user_id"] == "U-001"
    assert (test_settings.seed_dir / "T02_model_inventory.xlsx").exists()
    # Phase 4: seed evidence loads through the Document Centre batch path and feeds the documentation score
    doc_batches = client.get("/api/documents/batches", headers=login(client, *ADMIN)).json()
    assert doc_batches[0]["source"] == "seed" and doc_batches[0]["status"] == "Completed"
    assert doc_batches[0]["successful_records"] > 1000 and doc_batches[0]["failed_records"] == 0
    gaps = client.get("/api/documents/completeness", headers=login(client, *ADMIN), params={"missing_only": True}).json()
    assert 0 < len(gaps) < 256
    m12 = client.get("/api/models/M-0012", headers=login(client, *ADMIN)).json()
    doc_score = next(c for c in m12["score"]["components"] if c["key"] == "documentation")
    assert doc_score["score"] is not None


def test_reset_disabled(client, test_settings):
    test_settings.allow_reset = False
    resp = client.post("/api/admin/reset", headers=login(client, *ADMIN))
    assert resp.status_code == 403
